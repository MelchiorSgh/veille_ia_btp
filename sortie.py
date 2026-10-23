"""MODULE 9 - Sorties : fichier CSV cumulatif + rapport Markdown du jour."""
import csv, os
from datetime import datetime

CHAMPS = ["date_collecte", "date_pub", "source", "titre", "lien", "phase", "trl", "impact",
          "pertinence", "fiabilite", "fraicheur", "score", "methode", "resume_fr"]
NOMBRES = ["trl", "impact", "pertinence", "fiabilite", "fraicheur", "score"]


def ecrire_csv(chemin, lignes):
    nouveau = not os.path.exists(chemin)
    with open(chemin, "a", newline="", encoding="utf-8-sig") as f:
        ecrivain = csv.DictWriter(f, fieldnames=CHAMPS)
        if nouveau:
            ecrivain.writeheader()
        ecrivain.writerows(lignes)
    # Synchronisation automatique du classeur Excel (.xlsx)
    try:
        chemin_xlsx = os.path.splitext(chemin)[0] + ".xlsx"
        ecrire_excel(chemin_xlsx, lire_csv(chemin))
    except Exception as e:
        pass


def lire_csv(chemin):
    if not os.path.exists(chemin):
        return []
    with open(chemin, "r", encoding="utf-8-sig") as f:
        lignes = list(csv.DictReader(f))
    for l in lignes:
        for k in NOMBRES:
            try:
                l[k] = float(l[k])
            except (ValueError, TypeError):
                l[k] = 0.0
    return lignes


def lire_liens(chemin):
    """Tous les liens déjà traités (gardés OU écartés comme doublons), un par ligne."""
    if not os.path.exists(chemin):
        return set()
    with open(chemin, "r", encoding="utf-8") as f:
        return set(f.read().split())


def sauver_liens(chemin, liens):
    with open(chemin, "w", encoding="utf-8") as f:
        f.write("\n".join(sorted(liens)))


def ecrire_rapport(chemin, lignes, doublons, erreurs, seuil, stats, toutes_fiches=None):
    """Génère un rapport exécutif aéré, clair et structuré pour les décideurs et la direction."""
    date_str = f"{datetime.now():%d/%m/%Y}"
    fiches = lignes if lignes else (toutes_fiches or [])
    prioritaires = [f for f in fiches if float(f.get("score", 0)) >= seuil]
    
    def get_raison(item):
        if isinstance(item, tuple) and len(item) == 2:
            return str(item[1]).lower()
        if isinstance(item, dict):
            return str(item.get("motif_rejet", item.get("raison", ""))).lower()
        return str(item).lower()

    # Comptage des motifs de rejet
    nb_lexical = sum(1 for item in doublons if "lexical" in get_raison(item))
    nb_doublons = sum(1 for item in doublons if any(k in get_raison(item) for k in ("sens", "connu", "doublon")))
    nb_llm = sum(1 for item in doublons if any(k in get_raison(item) for k in ("llm", "hors sujet")))
    if not (nb_lexical or nb_doublons or nb_llm) and doublons:
        nb_lexical = len(doublons)

    L = [
        f"# Rapport Exécutif de Veille Technologique — {date_str}",
        "",
        "> **Dispositif de Veille :** Observatoire automatisé (IFRI - Licence IA)",
        "> **Thématique :** L'intelligence artificielle dans le BTP et l'architecture",
        "",
        "---",
        "",
        "## 1. Synthèse Décisionnelle & Indicateurs Clés",
        "",
        "| Indicateur Clé | Valeur | Description & Impact |",
        "| :--- | :---: | :--- |",
        f"| **Flux mondiaux surveillés (bruts)** | **{stats.get('lus', len(fiches))}** | Volume total d'articles collectés via les flux RSS/Atom |",
        f"| **Innovations validées en base** | **{len(fiches)}** | Fiches qualifiées par le double filtrage IA et stockées |",
        f"| **Signaux prioritaires (Score >= {seuil}/10)** | **{len(prioritaires)}** | Innovations majeures à fort impact opérationnel immédiat |",
        f"| **Articles écartés (Quarantaine)** | **{len(doublons)}** | Rejets conservés pour audit et vérification humaine |",
        "",
        "---",
        "",
        f"## 2. Top Signaux & Innovations Prioritaires (Score >= {seuil}/10)",
        ""
    ]

    if prioritaires:
        for p in prioritaires[:5]:
            L += [
                f"### [{p.get('titre')}]({p.get('lien')})",
                f"- **Phase :** {p.get('phase')} | **Source :** {p.get('source')} | **Score :** {float(p.get('score', 0)):.1f}/10 | **Maturité :** TRL {p.get('trl', '?')}/9",
                f"- **Synthèse :** {p.get('resume_fr', '')}",
                ""
            ]
    else:
        L += ["*Aucune nouvelle alerte prioritaire détectée sur ce lot spécifique.*", ""]

    L += [
        "---",
        "",
        "## 3. Répartition des Innovations par Cycle de Vie du Bâtiment",
        ""
    ]

    for phase, label in (("Conception", "Phase 1 : Conception & Modélisation BIM"),
                         ("Chantier", "Phase 2 : Chantier, Robotique & Sécurité Ouvrière"),
                         ("Exploitation", "Phase 3 : Exploitation, Maintenance & Jumeaux Numériques")):
        items = [x for x in fiches if x.get("phase") == phase]
        L += [f"### {label} ({len(items)} fiches)", ""]
        if items:
            for item in items[:4]:
                L += [f"- **[{item.get('titre')}]({item.get('lien')})** ({item.get('source')}, Score {float(item.get('score', 0)):.1f}/10, TRL {item.get('trl', '?')}) — *{item.get('resume_fr', '')[:140]}...*"]
            L += [""]
        else:
            L += ["*Aucune innovation répertoriée dans cette phase pour le moment.*", ""]

    L += [
        "---",
        "",
        "## 4. Bilan Qualité, Filtrage & Esprit Critique (Human-in-the-Loop)",
        "",
        "Afin de garantir l'absence d'hallucinations et de neutraliser les faux positifs, le pipeline applique un filtrage strict à 3 niveaux :",
        "",
        "| Motif d'exclusion | Nombre | Justification méthodologique |",
        "| :--- | :---: | :--- |",
        f"| **Filtre lexical (hors-sujet)** | {nb_lexical} | Absence des mots-clés bilingues du domaine dans le titre et le résumé |",
        f"| **Doublons sémantiques / URL** | {nb_doublons} | Articles déjà connus ou clones de communiqués de presse (similarité >= 85%) |",
        f"| **Refus sémantique LLM (Gemini)** | {nb_llm} | Mots-clés présents mais faux amis (ex: architecture informatique, campus sans IA) |",
        f"| **Total articles en quarantaine** | **{len(doublons)}** | Consultables et repêchables directement dans l'onglet Rejetés |",
        ""
    ]

    if erreurs:
        L += ["### Sources ayant rencontré une erreur technique :", ""] + [f"- {e}" for e in erreurs] + [""]

    L += [
        "---",
        "",
        "## 5. Recommandations Opérationnelles pour les Équipes",
        "",
        "1. **Accélération BIM/IA :** Prioriser l'expérimentation des agents d'automatisation des maquettes Revit/IFC pour réduire les temps de pré-construction.",
        "2. **Sécurité sur chantier :** Évaluer les systèmes de vision par ordinateur (détection de casques et EPI) pour réduire les risques d'accidents du travail.",
        "3. **Conformité & Éthique :** Veiller au respect du cadre légal (Code du numérique béninois et APDP) concernant le traitement d'images d'ouvriers sur chantier.",
        ""
    ]

    with open(chemin, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


CHAMPS_REJETS = ["date", "source", "titre", "lien", "motif_rejet", "decision_humaine"]


def ecrire_rejets(chemin, rejets):
    """Enregistre les articles rejetés sans doublons d'URL (pour Excel et vérification humaine)."""
    if not rejets:
        return
    lignes_existantes = []
    liens_vus = set()
    if os.path.exists(chemin):
        with open(chemin, "r", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                lien = row.get("lien", "").strip()
                if lien and lien not in liens_vus:
                    liens_vus.add(lien)
                    lignes_existantes.append(row)

    nouveaux = []
    maintenant_str = f"{datetime.now():%Y-%m-%d}"
    for item in rejets:
        if isinstance(item, tuple) and len(item) == 2:
            a, raison = item
        elif isinstance(item, dict):
            a, raison = item, item.get("motif_rejet", item.get("raison", ""))
        else:
            continue
        lien = a.get("lien", "").strip()
        if not lien or lien in liens_vus:
            continue
        liens_vus.add(lien)
        nouveaux.append({
            "date": maintenant_str,
            "source": a.get("source", "Inconnue"),
            "titre": a.get("titre", ""),
            "lien": lien,
            "motif_rejet": str(raison),
            "decision_humaine": "A vérifier (En attente)"
        })

    toutes_lignes = lignes_existantes + nouveaux
    with open(chemin, "w", newline="", encoding="utf-8-sig") as f:
        ecrivain = csv.DictWriter(f, fieldnames=CHAMPS_REJETS)
        ecrivain.writeheader()
        ecrivain.writerows(toutes_lignes)
    # Synchronisation automatique du classeur Excel (.xlsx)
    try:
        chemin_xlsx = os.path.splitext(chemin)[0] + ".xlsx"
        ecrire_excel_rejets(chemin_xlsx, toutes_lignes)
    except Exception as e:
        pass


def lire_rejets(chemin):
    if not os.path.exists(chemin):
        return []
    with open(chemin, "r", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def ecrire_excel(chemin_excel, lignes):
    """Génère un classeur Excel (.xlsx) professionnel, aéré, stylisé et filtrable."""
    if not lignes:
        return
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Innovations BTP & IA"
        ws.views.sheetView[0].showGridLines = True

        headers = [
            ("Date Collecte", 13),
            ("Date Pub", 13),
            ("Source", 20),
            ("Titre de l'innovation", 50),
            ("Phase BTP", 16),
            ("TRL", 8),
            ("Pertinence", 11),
            ("CRAAP", 10),
            ("Impact", 10),
            ("Fraîcheur", 10),
            ("Score /10", 12),
            ("Méthode", 12),
            ("Synthèse Opérationnelle", 65),
            ("Lien Web", 15)
        ]

        # En-têtes et largeurs
        for col_idx, (h_title, width) in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=h_title)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        center_align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_side = Side(border_style="thin", color="E2E8F0")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        ws.row_dimensions[1].height = 26
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align_header
            cell.border = cell_border

        font_regular = Font(name="Segoe UI", size=9.5)
        font_bold = Font(name="Segoe UI", size=9.5, bold=True)
        font_score_high = Font(name="Segoe UI", size=10, bold=True, color="166534")
        fill_score_high = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        align_left_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
        align_center = Alignment(horizontal="center", vertical="top")
        font_link = Font(name="Segoe UI", size=9.5, color="2563EB", underline="single")

        for row_idx, art in enumerate(lignes, 2):
            try:
                score_val = float(art.get("score", 0))
            except (ValueError, TypeError):
                score_val = 0.0

            is_zebra = (row_idx % 2 == 0)

            trl_val = art.get("trl", "")
            try:
                trl_val = int(float(trl_val))
            except (ValueError, TypeError):
                trl_val = str(trl_val)

            source_clean = art.get("source", "").replace("Google Actualités ", "Google Actu ")
            lien_url = art.get("lien", "").strip()

            vals = [
                (art.get("date_collecte", ""), align_center, None, font_regular, None),
                (art.get("date_pub", ""), align_center, None, font_regular, None),
                (source_clean, align_center, None, font_regular, None),
                (art.get("titre", ""), align_left_wrap, None, font_bold, None),
                (art.get("phase", ""), align_center, None, font_bold, None),
                (trl_val, align_center, None, font_regular, None),
                (float(art.get("pertinence", 0)), align_center, None, font_regular, "0.0"),
                (float(art.get("fiabilite", 0)), align_center, None, font_regular, "0.0"),
                (float(art.get("impact", 0)), align_center, None, font_regular, "0.0"),
                (float(art.get("fraicheur", 0)), align_center, None, font_regular, "0.0"),
                (score_val, align_center, fill_score_high if score_val >= 7.5 else None, font_score_high if score_val >= 7.5 else font_bold, "0.0"),
                (art.get("methode", ""), align_center, None, font_regular, None),
                (art.get("resume_fr", ""), align_left_wrap, None, font_regular, None),
                ("Consulter", align_center, None, font_link, None)
            ]

            for col_idx, (val, alignment, fill, font, num_format) in enumerate(vals, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font
                cell.alignment = alignment
                cell.border = cell_border
                if fill:
                    cell.fill = fill
                elif is_zebra:
                    cell.fill = fill_zebra

                if num_format:
                    cell.number_format = num_format

                if col_idx == 14 and lien_url.startswith("http"):
                    cell.hyperlink = lien_url

        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        wb.save(chemin_excel)
    except Exception as e:
        print(f"[!] Erreur génération Excel ({chemin_excel}) : {e}")


def ecrire_excel_rejets(chemin_excel, rejets):
    """Génère un classeur Excel (.xlsx) propre, aéré et lisible pour l'audit des articles en quarantaine."""
    if not rejets:
        return
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Rejets & Quarantaine"
        ws.views.sheetView[0].showGridLines = True

        headers = [
            ("Date", 13),
            ("Source", 20),
            ("Titre de l'article écarté", 55),
            ("Catégorie de rejet", 22),
            ("Motif détaillé d'exclusion", 50),
            ("Statut", 14),
            ("Lien Web", 15)
        ]

        # En-têtes et largeurs
        for col_idx, (h_title, width) in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=h_title)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        center_align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_side = Side(border_style="thin", color="E2E8F0")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        ws.row_dimensions[1].height = 26
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align_header
            cell.border = cell_border

        font_regular = Font(name="Segoe UI", size=9.5)
        font_bold = Font(name="Segoe UI", size=9.5, bold=True)
        font_link = Font(name="Segoe UI", size=9.5, color="2563EB", underline="single")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        fill_cat_lexical = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        fill_cat_doublon = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
        fill_cat_llm = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

        align_left_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
        align_center = Alignment(horizontal="center", vertical="top")

        for row_idx, r in enumerate(rejets, 2):
            is_zebra = (row_idx % 2 == 0)

            raw_motif = str(r.get("motif_rejet", "")).strip()
            if "lexical" in raw_motif.lower():
                categorie = "Filtre lexical"
                motif_court = "Mots-clés BTP / IA absents du titre et du résumé"
                fill_cat = fill_cat_lexical
            elif any(k in raw_motif.lower() for k in ("sens", "connu", "doublon")):
                categorie = "Doublon vectoriel"
                if "(" in raw_motif and ")" in raw_motif:
                    taux = raw_motif.split("(")[-1].split(")")[0]
                    motif_court = f"Similarité {taux} avec un article existant"
                else:
                    motif_court = "Article ou URL déjà enregistrée"
                fill_cat = fill_cat_doublon
            else:
                categorie = "Refus IA (Hors-sujet)"
                motif_court = raw_motif.replace("jugé hors sujet par le LLM : ", "").replace("jugé hors sujet par le LLM :", "").strip()
                fill_cat = fill_cat_llm

            source_clean = r.get("source", "").replace("Google Actualités ", "Google Actu ")
            decision = r.get("decision_humaine", "En attente")
            if "attente" in decision.lower():
                decision_clean = "En attente"
            elif "valid" in decision.lower():
                decision_clean = "Validé (Humain)"
            else:
                decision_clean = decision

            lien_url = r.get("lien", "").strip()

            vals = [
                (r.get("date", ""), align_center, None, font_regular),
                (source_clean, align_center, None, font_regular),
                (r.get("titre", ""), align_left_wrap, None, font_bold),
                (categorie, align_center, fill_cat, font_bold),
                (motif_court, align_left_wrap, None, font_regular),
                (decision_clean, align_center, None, font_regular),
                ("Consulter", align_center, None, font_link)
            ]

            for col_idx, (val, alignment, fill, font) in enumerate(vals, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font
                cell.alignment = alignment
                cell.border = cell_border
                if fill:
                    cell.fill = fill
                elif is_zebra:
                    cell.fill = fill_zebra

                if col_idx == 7 and lien_url.startswith("http"):
                    cell.hyperlink = lien_url

        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        wb.save(chemin_excel)
    except Exception as e:
        print(f"[!] Erreur génération Excel rejets ({chemin_excel}) : {e}")


